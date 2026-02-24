from __future__ import annotations

from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import numbers
from openpyxl.utils.dataframe import dataframe_to_rows

from market_report.observability.logging import get_logger


class ExcelWriter:
    """Write consolidated workbook with all tabs matching golden structure."""

    def __init__(self):
        self._logger = get_logger(self.__class__.__name__)

    def write_workbook(
        self,
        tabs: Dict[str, pd.DataFrame],
        output_path: Path,
        tab_order: list[str] | None = None,
    ) -> Path:
        """Write all tabs to Excel workbook.
        
        Args:
            tabs: Dictionary mapping tab names to DataFrames
            output_path: Path where to write the .xlsx file
            tab_order: Optional list specifying tab order (if None, uses tabs.keys())
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Use tab_order if provided, otherwise use natural order
        ordered_tabs = tab_order if tab_order else list(tabs.keys())

        self._logger.info("Writing workbook with %d tabs to %s", len(tabs), output_path)

        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            for tab_name in ordered_tabs:
                if tab_name not in tabs:
                    self._logger.warning("Tab '%s' not found in tabs dictionary, skipping", tab_name)
                    continue
                
                df = tabs[tab_name]
                if df.empty:
                    self._logger.warning("Tab '%s' is empty, writing empty sheet", tab_name)
                    # Write empty DataFrame to maintain structure
                    pd.DataFrame().to_excel(writer, sheet_name=tab_name, index=False)
                else:
                    df.to_excel(writer, sheet_name=tab_name, index=False, header=True)
                    self._logger.debug("Wrote tab '%s' with %d rows x %d cols", tab_name, len(df), len(df.columns))

        # Apply number formats for specific tabs after writing
        try:
            workbook = load_workbook(output_path)
            if "Histórico_Mercado" in workbook.sheetnames:
                sheet = workbook["Histórico_Mercado"]
                header_row = [cell.value for cell in sheet[1]]

                def _find_col(name: str) -> int | None:
                    for idx, value in enumerate(header_row, start=1):
                        if value == name:
                            return idx
                    return None

                price_col = _find_col("Precio Promedio Inv.")
                xm2_col = _find_col("$M2 Promedio Inv")

                currency_format = '"$"#,##0'
                if price_col:
                    for row in sheet.iter_rows(min_row=2, min_col=price_col, max_col=price_col):
                        for cell in row:
                            if isinstance(cell.value, (int, float)):
                                cell.number_format = currency_format
                if xm2_col:
                    for row in sheet.iter_rows(min_row=2, min_col=xm2_col, max_col=xm2_col):
                        for cell in row:
                            if isinstance(cell.value, (int, float)):
                                cell.number_format = currency_format

            workbook.save(output_path)
        except Exception as exc:
            self._logger.warning("Failed applying number formats: %s", exc)

        self._logger.info("Workbook written successfully: %s", output_path)
        return output_path

    def export_tabs_to_csv(
        self, tabs: Dict[str, pd.DataFrame], output_dir: Path
    ) -> Dict[str, Path]:
        """Export each tab to CSV for auditing.
        
        Returns:
            Dictionary mapping tab names to CSV file paths
        """
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)

        csv_paths: Dict[str, Path] = {}
        for tab_name, df in tabs.items():
            csv_path = output_dir / f"{tab_name}.csv"
            df.to_csv(csv_path, index=False, header=False, encoding="utf-8")
            csv_paths[tab_name] = csv_path
            self._logger.debug("Exported tab '%s' to CSV: %s", tab_name, csv_path)

        return csv_paths
