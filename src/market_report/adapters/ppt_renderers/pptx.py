from __future__ import annotations

from pathlib import Path
from typing import Dict

from pptx import Presentation
from pptx.util import Inches

from market_report.config.schemas import PptConfig
from market_report.ports.ppt import PptRenderer


class PptxRenderer(PptRenderer):
    """Render PPTX using python-pptx."""

    def render(self, ppt_config: PptConfig, chart_paths: Dict[str, Path], output_path: Path) -> Path:
        presentation = self._load_template(ppt_config.template)
        layout = presentation.slide_layouts[5] if presentation.slide_layouts else presentation.slide_layouts[0]

        for slide_cfg in ppt_config.slides:
            slide = presentation.slides.add_slide(layout)
            self._add_title(slide, slide_cfg.title)
            chart_path = chart_paths.get(slide_cfg.chart)
            if chart_path and chart_path.exists():
                slide.shapes.add_picture(str(chart_path), Inches(1), Inches(1.8), width=Inches(8))

        presentation.save(str(output_path))
        return output_path

    def _load_template(self, template_path: str) -> Presentation:
        path = Path(template_path)
        if path.exists():
            try:
                return Presentation(str(path))
            except Exception:
                pass
        return Presentation()

    def _add_title(self, slide, title: str) -> None:
        if slide.shapes.title:
            slide.shapes.title.text = title
        else:
            textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
            textbox.text_frame.text = title
